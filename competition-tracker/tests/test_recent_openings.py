from datetime import date

import openpyxl

from recent_openings import (
    OTHER_OPENINGS_COLUMNS,
    OTHER_OPENINGS_SHEET,
    RECENT_OPENINGS_COLUMNS,
    RECENT_OPENINGS_SHEET,
    TO_VERIFY_COLUMNS,
    TO_VERIFY_SHEET,
    build_recent_openings_rows,
    write_recent_openings_workbook,
)


def _brand_result(brand="Maison Francis Kurkdjian", slug="mfk", stores=None, status="ok"):
    return {
        "brand": brand,
        "slug": slug,
        "status": status,
        "stores": stores or [],
    }


def _store(name="Maison Francis Kurkdjian Houston", address="4444 Westheimer Rd", city="Houston", country="United States"):
    return {
        "name": name,
        "address": address,
        "city": city,
        "country": country,
        "region": "NOAM",
    }


def _entry(**overrides):
    base = {
        "brand": "Maison Francis Kurkdjian",
        "region": "NOAM",
        "country": "United States",
        "city": "Houston",
        "store_name": "Maison Francis Kurkdjian Houston",
        "address": "4444 Westheimer Rd",
        "url": "https://example.com/opening",
        "source_title": "Opening News",
        "source_domain": "example.com",
        "source_type": "industry_publication",
        "source_date": "2026-07-01",
        "store_classification": "FSS",
        "confidence": "medium",
        "status": "CONFIRMED",
        "verification_notes": "Verified",
        "evidence": "Evidence",
    }
    base.update(overrides)
    return base


def _run(results, bible_index, cache_entries, **kwargs):
    kwargs.setdefault("checked_date", date(2026, 7, 17))
    return build_recent_openings_rows(results, bible_index, {"entries": cache_entries}, **kwargs)


def test_store_already_in_bible_is_excluded():
    results = [_brand_result(stores=[_store()])]
    bible_index = {
        "by_door_key": {
            ("maison francis kurkdjian", "united states", "houston", "maison francis kurkdjian houston"): [{}],
        }
    }
    recent_rows, verify_rows, other_rows, summary = _run(results, bible_index, [_entry()])
    assert recent_rows == []
    assert verify_rows == []
    assert other_rows == []
    assert summary.recent_openings_found == 0
    assert summary.excluded_in_bible == 1


def test_scraper_door_name_differs_from_bible_but_entry_name_matches_it():
    """Mirrors Parfums de Marly Le Marais: the live locator scrape names the
    door differently ("Boutique Marais") than the BIBLE does ("LE MARAIS"),
    but the research entry's own store_name matches the BIBLE exactly. The
    pair must be recognized as already tracked either way, not reported as
    a new opening just because the scraper's own wording didn't match."""
    results = [_brand_result(
        brand="Parfums de Marly", slug="pdm",
        stores=[_store(
            name="Parfums de Marly, Boutique Marais",
            address="45, rue Vieille du Temple",
            city="Paris", country="France",
        )],
    )]
    bible_index = {
        "by_door_key": {
            ("parfums de marly", "france", "paris", "parfums de marly le marais"): [{}],
        }
    }
    entry = _entry(
        brand="Parfums de Marly", region="EMEA", country="France", city="Paris",
        store_name="Parfums de Marly Le Marais",
        address="45 rue Vieille du Temple, 75004 Paris, France",
        source_type="official_brand_store_page", source_date="2026-06-20",
    )
    recent_rows, verify_rows, other_rows, summary = _run(results, bible_index, [entry])
    assert recent_rows == []
    assert verify_rows == []
    assert other_rows == []
    assert summary.excluded_in_bible == 1


def test_department_store_opening_goes_to_other_openings():
    recent_rows, verify_rows, other_rows, summary = _run(
        [_brand_result(stores=[_store()])],
        {"by_door_key": {}},
        [_entry(store_classification="DEPARTMENT_STORE")],
    )
    assert recent_rows == []
    assert verify_rows == []
    assert len(other_rows) == 1
    assert other_rows[0]["DOOR TYPE"] == "DEPARTMENT_STORE"
    assert summary.other_openings_found == 1
    assert summary.other_openings_by_classification == {"DEPARTMENT_STORE": 1}


def test_travel_retail_opening_goes_to_other_openings_only_when_flag_is_set():
    entry = [_entry(store_classification="TRAVEL_RETAIL_BOUTIQUE")]

    off_by_default = _run([_brand_result(stores=[_store()])], {"by_door_key": {}}, entry)
    assert off_by_default[2] == []  # other_rows
    assert off_by_default[3].excluded_travel_retail_gated == 1

    with_flag = _run(
        [_brand_result(stores=[_store()])], {"by_door_key": {}}, entry, include_travel_retail=True,
    )
    assert len(with_flag[2]) == 1
    assert with_flag[2][0]["DOOR TYPE"] == "TRAVEL_RETAIL_BOUTIQUE"
    assert with_flag[3].excluded_travel_retail_gated == 0


def test_perfumery_and_multibrand_retailer_are_excluded_not_other_openings():
    for classification in ("PERFUMERY", "MULTIBRAND_RETAILER", "ONLINE"):
        recent_rows, verify_rows, other_rows, summary = _run(
            [_brand_result(stores=[_store()])],
            {"by_door_key": {}},
            [_entry(store_classification=classification)],
        )
        assert recent_rows == verify_rows == other_rows == []
        assert summary.excluded_non_fss_fsf == 1


def test_opening_source_within_60_days_is_included():
    recent_rows, verify_rows, other_rows, summary = _run(
        [_brand_result(stores=[_store()])],
        {"by_door_key": {}},
        [_entry(source_date="2026-07-10")],
    )
    assert len(recent_rows) == 1
    assert verify_rows == []
    assert other_rows == []
    assert recent_rows[0]["OPENING DATE"] == "2026-07-10"
    assert summary.recent_openings_found == 1


def test_old_announcement_is_excluded_not_to_verify():
    """A known opening date outside the recent window is a resolved fact,
    not an open question - it must be excluded outright, never parked in
    TO VERIFY (TO VERIFY is reserved for unknown-date candidates only)."""
    recent_rows, verify_rows, other_rows, summary = _run(
        [_brand_result(stores=[_store()])],
        {"by_door_key": {}},
        [_entry(source_date="2026-04-01")],
    )
    assert recent_rows == []
    assert verify_rows == []
    assert other_rows == []
    assert summary.to_verify_count == 0
    assert summary.excluded_old_openings == 1


def test_known_old_date_from_a_source_below_priority_threshold_is_still_excluded_not_to_verify():
    """Even when the source itself is too weak to promote to RECENT
    OPENINGS, a known-but-old date must not be reported as if the date
    were unknown."""
    recent_rows, verify_rows, other_rows, summary = _run(
        [_brand_result(stores=[_store()])],
        {"by_door_key": {}},
        [_entry(source_date="2026-04-01", source_type="unlisted_source_type")],
    )
    assert recent_rows == []
    assert verify_rows == []
    assert other_rows == []
    assert summary.excluded_old_openings == 1


def test_unknown_opening_date_goes_to_to_verify():
    recent_rows, verify_rows, other_rows, summary = _run(
        [_brand_result(stores=[_store()])],
        {"by_door_key": {}},
        [_entry(source_date=None)],
    )
    assert recent_rows == []
    assert len(verify_rows) == 1
    assert verify_rows[0]["REASON TO VERIFY"] == "opening source has no date"
    assert verify_rows[0]["POSSIBLE DOOR TYPE"] == "FSS"
    assert other_rows == []
    assert summary.to_verify_count == 1


def test_duplicate_article_and_locator_result_merge_into_one_row():
    recent_rows, verify_rows, other_rows, summary = _run(
        [_brand_result(stores=[_store(), _store()])],
        {"by_door_key": {}},
        [_entry(), _entry(source_title="Backup Source")],
    )
    assert len(recent_rows) == 1
    assert verify_rows == []
    assert other_rows == []
    assert summary.recent_openings_found == 1


def test_recent_official_announcement_absent_from_bible_is_included_without_store_locator_match():
    results = [_brand_result(brand="Byredo", slug="byredo", stores=[])]
    recent_rows, verify_rows, other_rows, summary = _run(
        results,
        {"by_door_key": {}},
        [_entry(
            brand="Byredo",
            region="APAC",
            country="Hong Kong",
            city="Hong Kong",
            store_name="Byredo Gough Street",
            address="2-10 Gough Street, Hong Kong",
            url="https://example.com/byredo",
            source_title="Byredo Opens New Boutique",
            source_type="official_brand_newsroom",
            source_date="2026-07-10",
            store_classification="FSS",
            confidence="high",
            verification_notes="Official opening post",
        )],
    )
    assert len(recent_rows) == 1
    assert verify_rows == []
    assert other_rows == []
    assert recent_rows[0]["BRAND"] == "Byredo"
    assert summary.recent_openings_by_brand == {"Byredo": 1}


def test_credible_mall_or_landlord_announcement_is_included():
    results = [_brand_result(stores=[])]
    recent_rows, verify_rows, other_rows, summary = _run(
        results,
        {"by_door_key": {}},
        [_entry(
            city="Miami",
            store_name="Maison Francis Kurkdjian - Miami Design District",
            address="176 NE 41st St, Miami, FL",
            url="https://example.com/mall",
            source_title="Miami Design District",
            source_type="official_mall_landlord_directory",
            source_date="2026-07-05",
            store_classification="FSS",
            confidence="high",
            verification_notes="Official landlord directory",
        )],
    )
    assert len(recent_rows) == 1
    assert verify_rows == []
    assert summary.recent_openings_found == 1


def test_official_source_outranks_industry_news_when_both_present():
    """Two independent sources agree on the same door - the official one
    must be chosen as SOURCE, the industry one recorded as SECOND SOURCE."""
    official = _entry(
        url="https://brand.example/store", source_title="Brand Official Store Page",
        source_type="official_brand_store_page", source_date="2026-07-10",
    )
    industry = _entry(
        url="https://fashionnetwork.example/article", source_title="FashionNetwork Coverage",
        source_type="industry_publication", source_date="2026-07-09",
    )
    recent_rows, verify_rows, other_rows, summary = _run(
        [_brand_result(stores=[_store()])], {"by_door_key": {}}, [industry, official],
    )
    assert len(recent_rows) == 1
    assert recent_rows[0]["SOURCE"] == "Brand Official Store Page"
    assert recent_rows[0]["SECOND SOURCE"] == "FashionNetwork Coverage"
    assert recent_rows[0]["SECOND SOURCE URL"] == "https://fashionnetwork.example/article"


def test_curated_industry_news_source_excluded_by_default_but_included_with_flag():
    """A finding sourced only from a curated news outlet (news_sources.yaml)
    must not surface at all by default, but can discover an opening the
    store locator never returned once --include-industry-news is passed."""
    entry = [_entry(
        source_type="fashionnetwork", source_title="FashionNetwork", source_date="2026-07-10",
    )]
    results = [_brand_result(stores=[])]

    off_by_default = _run(results, {"by_door_key": {}}, entry)
    assert off_by_default[0] == []  # recent_rows
    assert off_by_default[3].excluded_industry_news_gated == 1

    with_flag = _run(results, {"by_door_key": {}}, entry, include_industry_news=True)
    assert len(with_flag[0]) == 1
    assert with_flag[3].excluded_industry_news_gated == 0


def test_linkedin_post_without_original_source_goes_to_to_verify():
    results = [_brand_result(stores=[])]
    entry = [_entry(
        source_type="linkedin_post",
        source_title="Some LinkedIn Account",
        url="https://linkedin.com/posts/example",
        source_date="2026-07-10",
        status="CONFIRMED",
    )]
    recent_rows, verify_rows, other_rows, summary = _run(results, {"by_door_key": {}}, entry)
    assert recent_rows == []
    assert other_rows == []
    assert len(verify_rows) == 1
    assert "LinkedIn" in verify_rows[0]["REASON TO VERIFY"]


def test_linkedin_post_with_original_source_can_reach_recent_openings():
    results = [_brand_result(stores=[])]
    entry = [_entry(
        source_type="linkedin_post",
        source_title="Some LinkedIn Account",
        url="https://linkedin.com/posts/example",
        original_source_url="https://brand.example/newsroom/opening",
        source_date="2026-07-10",
        status="CONFIRMED",
    )]
    recent_rows, verify_rows, other_rows, summary = _run(results, {"by_door_key": {}}, entry)
    assert len(recent_rows) == 1
    assert verify_rows == []


def test_workbook_contains_all_three_sheets_with_required_columns(tmp_path):
    output_path = tmp_path / "recent_openings.xlsx"
    write_recent_openings_workbook(
        [{col: "x" for col in RECENT_OPENINGS_COLUMNS}],
        [{col: "x" for col in TO_VERIFY_COLUMNS}],
        [{col: "x" for col in OTHER_OPENINGS_COLUMNS}],
        output_path,
    )

    workbook = openpyxl.load_workbook(output_path)
    try:
        assert workbook.sheetnames == [RECENT_OPENINGS_SHEET, TO_VERIFY_SHEET, OTHER_OPENINGS_SHEET]
        recent_sheet = workbook[RECENT_OPENINGS_SHEET]
        verify_sheet = workbook[TO_VERIFY_SHEET]
        other_sheet = workbook[OTHER_OPENINGS_SHEET]
        assert [cell.value for cell in recent_sheet[1]] == RECENT_OPENINGS_COLUMNS
        assert [cell.value for cell in verify_sheet[1]] == TO_VERIFY_COLUMNS
        assert [cell.value for cell in other_sheet[1]] == OTHER_OPENINGS_COLUMNS
        assert recent_sheet.freeze_panes == "A2"
        assert verify_sheet.freeze_panes == "A2"
        assert other_sheet.freeze_panes == "A2"
        assert recent_sheet.auto_filter.ref == recent_sheet.dimensions
    finally:
        workbook.close()


def test_hyperlinks_are_clickable(tmp_path):
    output_path = tmp_path / "recent_openings.xlsx"
    recent_row = {col: "" for col in RECENT_OPENINGS_COLUMNS}
    recent_row["SOURCE URL"] = "https://example.com/opening"
    recent_row["SECOND SOURCE URL"] = "https://example.com/second"
    write_recent_openings_workbook([recent_row], [], [], output_path)

    workbook = openpyxl.load_workbook(output_path)
    try:
        sheet = workbook[RECENT_OPENINGS_SHEET]
        source_col = RECENT_OPENINGS_COLUMNS.index("SOURCE URL") + 1
        second_col = RECENT_OPENINGS_COLUMNS.index("SECOND SOURCE URL") + 1
        source_cell = sheet.cell(row=2, column=source_col)
        second_cell = sheet.cell(row=2, column=second_col)
        assert source_cell.hyperlink is not None
        assert source_cell.hyperlink.target == "https://example.com/opening"
        assert second_cell.hyperlink is not None
        assert second_cell.hyperlink.target == "https://example.com/second"
    finally:
        workbook.close()
