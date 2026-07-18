"""Etape (recent openings): surfaces confirmed newly-opened FSS/FSF doors so
they can be added to the BIBLE, using data/online_research_cache.json as the
evidence source (see online_research.py's module docstring for why that's a
curated cache rather than a live search) - now broadened to also draw on
curated industry/luxury/beauty/travel-retail press and LinkedIn (see
news_sources.yaml) alongside official brand/mall/landlord sources.

Two independent discovery routes feed the same check - neither gates the
other, and a finding never has to come from both:

  A. store-locator candidate -> recent-opening evidence: the scraper already
     returned this door (results[i]["stores"]) and online research
     separately dates its opening.
  B. recent-opening announcement -> official locator + BIBLE check: a dated,
     credible cache entry, whether or not the scraper's own coverage
     happened to include the same door (a fixed seed list, a paginated
     locator, a blocked request, or a city-name mismatch between sources can
     all cause the scraper to miss a real door - that is not a reason to
     miss the opening too).

Every cache entry gets classified into exactly one bucket:

  - RECENT OPENINGS: FSS/FSF, CONFIRMED/PROBABLE, dated within the recent
    window, not already in the BIBLE, from a recognized source, and (if the
    only evidence is an unattributed LinkedIn repost) still routed to TO
    VERIFY instead - see _is_unattributed_linkedin_repost.
  - TO VERIFY: FSS/FSF but missing what RECENT OPENINGS requires (usually a
    dated source, or - for LinkedIn - an identified original source).
  - OTHER OPENINGS: a real, non-FSS/FSF retail development worth knowing
    about (travel-retail boutique, department-store opening, corner/
    concession, shop-in-shop, pop-up, relocation, reopening, or an
    otherwise-unclear non-FSS development) - reported separately, never
    mixed into RECENT OPENINGS.
  - excluded entirely: already in the BIBLE; a known opening date outside
    the recent window (a resolved fact, not an open question - never left
    in TO VERIFY); or a listing that isn't the brand's own doing at all
    (PERFUMERY / MULTIBRAND_RETAILER / ONLINE - a reseller carrying the
    brand, not the brand opening something).

--include-industry-news and --include-travel-retail (see run_report.py)
gate the newly-added curated press sources and travel-retail findings
respectively, off by default so the ordinary run stays conservative - see
CURATED_NEWS_SOURCE_TYPES / the TRAVEL_RETAIL_BOUTIQUE check below.
"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from brand_aliases import brand_match_key
from diff_existing import _door_key, build_bible_index, filter_bible_rows, load_bible_competition
from news_sources import source_priorities as _load_news_source_priorities
from normalize import normalize
from online_research import CACHE_PATH_DEFAULT, load_cache

RECENT_OPENINGS_SHEET = "RECENT OPENINGS"
TO_VERIFY_SHEET = "TO VERIFY"
OTHER_OPENINGS_SHEET = "OTHER OPENINGS"

RECENT_OPENINGS_COLUMNS = [
    "BRAND",
    "REGION",
    "COUNTRY",
    "CITY",
    "DOOR NAME",
    "FULL ADDRESS",
    "DOOR TYPE",
    "OPENING DATE",
    "SOURCE",
    "SOURCE URL",
    "SECOND SOURCE",
    "SECOND SOURCE URL",
    "DATE CHECKED",
    "NOTES",
]

TO_VERIFY_COLUMNS = [
    "BRAND",
    "REGION",
    "COUNTRY",
    "CITY",
    "DOOR NAME",
    "FULL ADDRESS",
    "POSSIBLE DOOR TYPE",
    "REASON TO VERIFY",
    "SOURCE",
    "SOURCE URL",
    "SECOND SOURCE",
    "SECOND SOURCE URL",
    "DATE CHECKED",
    "NOTES",
]

OTHER_OPENINGS_COLUMNS = [
    "BRAND",
    "REGION",
    "COUNTRY",
    "CITY",
    "DOOR NAME",
    "FULL ADDRESS",
    "DOOR TYPE",
    "OPENING DATE",
    "SOURCE",
    "SOURCE URL",
    "DATE CHECKED",
    "NOTES",
]

# Simple, deterministic 3-tier SOURCE_PRIORITY (see news_sources.yaml's
# header for the exact tier definitions) - tier 1 for anything official
# (brand or mall/landlord), tier 2-3 merged in from news_sources.yaml so
# that file stays the single source of truth for the curated press tiers.
BASE_SOURCE_PRIORITY = {
    "official_brand_website": 1,
    "official_brand_newsroom": 1,
    "official_brand_store_page": 1,
    "official_verified_brand_social_post": 1,  # the brand's own official LinkedIn
    "official_mall_landlord_directory": 1,
    "official_mall_landlord_website": 1,
    "official_airport_website": 1,
    "official_shopping_centre_website": 1,
    "industry_publication": 2,
    "beauty_publication": 2,
    "luxury_publication": 2,
    "retail_publication": 2,
    "travel_retail_publication": 2,
    "local_publication": 3,
    # A LinkedIn post that is NOT the brand/mall's own official page (an
    # unverified account resharing news) - see _is_unattributed_linkedin_repost.
    "linkedin_post": 2,
}
RECENT_SOURCE_TYPES = {**BASE_SOURCE_PRIORITY, **_load_news_source_priorities()}

# source_type slugs sourced from news_sources.yaml (FashionNetwork, Moodie
# Davitt, TRBusiness, ...) - only considered at all when --include-industry-
# news is passed, keeping the default run conservative (see module docstring).
CURATED_NEWS_SOURCE_TYPES = frozenset(_load_news_source_priorities().keys())

# Legitimate non-FSS/FSF retail developments worth reporting separately
# (OTHER OPENINGS) - as opposed to PERFUMERY/MULTIBRAND_RETAILER/ONLINE,
# which mean "not the brand's own doing" and are excluded outright.
OTHER_OPENING_CLASSIFICATIONS = {
    "TRAVEL_RETAIL_BOUTIQUE", "DEPARTMENT_STORE", "CORNER_OR_CONCESSION",
    "SHOP_IN_SHOP", "POP_UP", "RELOCATION", "REOPENING", "UNCLEAR",
}

ALLOWED_DOOR_TYPES = {"FSS", "FSF"}
RECENT_OPENING_STATUSES = {"CONFIRMED", "PROBABLE"}
VERIFYABLE_STATUSES = {"CONFIRMED", "PROBABLE", "TO VERIFY"}
ROW_KEY_FIELDS = ("BRAND", "COUNTRY", "CITY", "DOOR NAME")


@dataclass(frozen=True)
class PipelineSummary:
    brands_checked: int
    recent_openings_found: int
    to_verify_count: int
    excluded_non_fss_fsf: int
    excluded_in_bible: int
    excluded_old_openings: int = 0
    other_openings_found: int = 0
    excluded_travel_retail_gated: int = 0
    excluded_industry_news_gated: int = 0
    recent_openings_by_brand: dict[str, int] = field(default_factory=dict)
    to_verify_by_brand: dict[str, int] = field(default_factory=dict)
    other_openings_by_brand: dict[str, int] = field(default_factory=dict)
    other_openings_by_classification: dict[str, int] = field(default_factory=dict)
    source_links_used: tuple[str, ...] = ()


def _parse_source_date(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except ValueError:
        return None


def _is_within_recent_days(opening_date: date | None, checked_date: date, recent_days: int) -> bool:
    if opening_date is None:
        return False
    earliest = checked_date - timedelta(days=recent_days)
    return earliest <= opening_date <= checked_date


def _source_priority(entry: dict) -> int:
    return RECENT_SOURCE_TYPES.get(entry.get("source_type"), 99)


def _is_unattributed_linkedin_repost(entry: dict) -> bool:
    """A LinkedIn post that isn't the brand/mall's own official page (see
    "official_verified_brand_social_post" above, which IS official) and
    doesn't preserve the original announcement it's reposting - never
    promoted to RECENT OPENINGS on its own, always routed to TO VERIFY
    instead, regardless of how confident/dated it otherwise looks."""
    return entry.get("source_type") == "linkedin_post" and not entry.get("original_source_url")


def _store_lookup_key(brand: str, store: dict) -> tuple[str, str, str, str]:
    return _door_key(brand, store.get("country"), store.get("city"), store.get("name"))


def _brandless_tokens(text: str | None, brand: str) -> set[str]:
    tokens = normalize(text).split()
    brand_tokens = set(normalize(brand).split())
    return {token for token in tokens if token not in brand_tokens}


def _entry_matches_store(entry: dict, brand_name: str, store: dict) -> bool:
    if brand_match_key(entry.get("brand")) != brand_match_key(brand_name):
        return False

    store_country = normalize(store.get("country"))
    entry_country = normalize(entry.get("country"))
    if store_country and entry_country and store_country != entry_country:
        return False

    store_city = normalize(store.get("city"))
    entry_city = normalize(entry.get("city"))
    if store_city and entry_city and store_city != entry_city:
        return False

    store_name = normalize(store.get("name"))
    entry_name = normalize(entry.get("store_name"))
    store_address = normalize(store.get("address"))
    entry_address = normalize(entry.get("address"))

    if store_name and entry_name and store_name == entry_name:
        return True
    if store_address and entry_address and store_address == entry_address:
        return True
    if store_address and entry_address and (store_address in entry_address or entry_address in store_address):
        return True

    store_tokens = _brandless_tokens(store.get("name"), brand_name)
    entry_tokens = _brandless_tokens(entry.get("store_name"), brand_name)
    if store_city and entry_city and store_tokens and entry_tokens and store_tokens.intersection(entry_tokens):
        return True

    return False


def _best_entry(entries: list[dict]) -> dict | None:
    if not entries:
        return None
    return sorted(
        entries,
        key=lambda entry: (
            _source_priority(entry),
            0 if _parse_source_date(entry.get("source_date")) else 1,
            -int(entry.get("status") == "CONFIRMED"),
            entry.get("source_title") or "",
        ),
    )[0]


def _recent_sort_key(entry: dict) -> tuple[int, int, str]:
    opening_date = _parse_source_date(entry.get("source_date"))
    return (
        _source_priority(entry),
        -opening_date.toordinal() if opening_date else 0,
        entry.get("source_title") or "",
    )


def _second_source_entry(matched_entries: list[dict], primary: dict) -> dict | None:
    """The next-best entry with a genuinely different URL than the primary
    source - simple deterministic "do we have independent corroboration"
    signal, no scoring formula (see EVIDENCE RULES' "preferred confirmation"
    - this is informational, never a hard requirement)."""
    primary_url = primary.get("url")
    for entry in sorted(matched_entries, key=_recent_sort_key):
        if entry is primary:
            continue
        if entry.get("url") and entry.get("url") != primary_url:
            return entry
    return None


def load_bible_index(existing_file: Path | str) -> dict:
    rows = load_bible_competition(existing_file)
    filtered_rows, _ = filter_bible_rows(rows)
    return build_bible_index(filtered_rows)


def load_research_cache(cache_path: Path | str = CACHE_PATH_DEFAULT) -> dict:
    return load_cache(cache_path)


def _format_source_label(entry: dict) -> str:
    return entry.get("source_title") or entry.get("source_domain") or "Source"


def _recent_row(brand_name: str, store: dict, entry: dict, matched_entries: list[dict], checked_date: date) -> dict:
    notes = entry.get("verification_notes") or ""
    if entry.get("status") == "PROBABLE":
        notes = "PROBABLE - " + notes if notes else "PROBABLE"
    second = _second_source_entry(matched_entries, entry)
    return {
        "BRAND": brand_name,
        "REGION": store.get("region") or entry.get("region"),
        "COUNTRY": store.get("country") or entry.get("country"),
        "CITY": store.get("city") or entry.get("city"),
        "DOOR NAME": store.get("name") or entry.get("store_name"),
        "FULL ADDRESS": store.get("address") or entry.get("address"),
        "DOOR TYPE": entry.get("store_classification"),
        "OPENING DATE": entry.get("source_date"),
        "SOURCE": _format_source_label(entry),
        "SOURCE URL": entry.get("url"),
        "SECOND SOURCE": _format_source_label(second) if second else "",
        "SECOND SOURCE URL": second.get("url") if second else "",
        "DATE CHECKED": checked_date.isoformat(),
        "NOTES": notes,
    }


def _verify_row(
    brand_name: str, store: dict, entry: dict, matched_entries: list[dict],
    checked_date: date, reason: str, possible_door_type: str | None,
) -> dict:
    second = _second_source_entry(matched_entries, entry)
    return {
        "BRAND": brand_name,
        "REGION": store.get("region") or entry.get("region"),
        "COUNTRY": store.get("country") or entry.get("country"),
        "CITY": store.get("city") or entry.get("city"),
        "DOOR NAME": store.get("name") or entry.get("store_name"),
        "FULL ADDRESS": store.get("address") or entry.get("address"),
        "POSSIBLE DOOR TYPE": possible_door_type,
        "REASON TO VERIFY": reason,
        "SOURCE": _format_source_label(entry),
        "SOURCE URL": entry.get("url"),
        "SECOND SOURCE": _format_source_label(second) if second else "",
        "SECOND SOURCE URL": second.get("url") if second else "",
        "DATE CHECKED": checked_date.isoformat(),
        "NOTES": entry.get("verification_notes") or entry.get("evidence") or "",
    }


def _other_row(brand_name: str, store: dict, entry: dict, checked_date: date) -> dict:
    return {
        "BRAND": brand_name,
        "REGION": store.get("region") or entry.get("region"),
        "COUNTRY": store.get("country") or entry.get("country"),
        "CITY": store.get("city") or entry.get("city"),
        "DOOR NAME": store.get("name") or entry.get("store_name"),
        "FULL ADDRESS": store.get("address") or entry.get("address"),
        "DOOR TYPE": entry.get("store_classification"),
        "OPENING DATE": entry.get("source_date"),
        "SOURCE": _format_source_label(entry),
        "SOURCE URL": entry.get("url"),
        "DATE CHECKED": checked_date.isoformat(),
        "NOTES": entry.get("verification_notes") or entry.get("evidence") or "",
    }


def _row_key(row: dict) -> tuple[str, ...]:
    return tuple(normalize(row.get(field)) for field in ROW_KEY_FIELDS)


def _store_from_entry(entry: dict) -> dict:
    return {
        "name": entry.get("store_name"),
        "address": entry.get("address"),
        "city": entry.get("city"),
        "country": entry.get("country"),
        "region": entry.get("region"),
    }


def _candidate_reason(best_entry: dict) -> str | None:
    """Only ever called once every matched entry's opening date is unknown
    (see process_candidate) - so a known-but-old date is never described
    here as "to verify"; it's excluded entirely instead."""
    if _source_priority(best_entry) >= 99:
        return "source type is below the accepted priority threshold"
    return "opening source has no date"


def _selected_brand_map(results: list[dict]) -> dict[str, str]:
    selected = {}
    for result in results:
        brand_name = result.get("brand")
        if brand_name:
            selected[brand_match_key(brand_name)] = brand_name
    return selected


def build_recent_openings_rows(
    results: list[dict],
    bible_index: dict,
    research_cache: dict,
    recent_days: int = 60,
    checked_date: date | None = None,
    include_industry_news: bool = False,
    include_travel_retail: bool = False,
) -> tuple[list[dict], list[dict], list[dict], PipelineSummary]:
    checked_date = checked_date or date.today()
    recent_rows: list[dict] = []
    verify_rows: list[dict] = []
    other_rows: list[dict] = []
    recent_row_keys: set[tuple[str, ...]] = set()
    verify_row_keys: set[tuple[str, ...]] = set()
    other_row_keys: set[tuple[str, ...]] = set()
    recent_by_brand: defaultdict[str, int] = defaultdict(int)
    verify_by_brand: defaultdict[str, int] = defaultdict(int)
    other_by_brand: defaultdict[str, int] = defaultdict(int)
    other_by_classification: defaultdict[str, int] = defaultdict(int)
    source_links_used: set[str] = set()
    excluded_non_fss_fsf = 0
    excluded_in_bible = 0
    excluded_old_openings = 0
    excluded_travel_retail_gated = 0
    excluded_industry_news_gated = 0

    selected_brands = _selected_brand_map(results)
    candidate_entries = [
        entry for entry in research_cache.get("entries", [])
        if brand_match_key(entry.get("brand")) in selected_brands
    ]

    entries = []
    for entry in candidate_entries:
        if entry.get("store_classification") == "TRAVEL_RETAIL_BOUTIQUE" and not include_travel_retail:
            excluded_travel_retail_gated += 1
            continue
        if entry.get("source_type") in CURATED_NEWS_SOURCE_TYPES and not include_industry_news:
            excluded_industry_news_gated += 1
            continue
        entries.append(entry)

    processed_candidates: set[tuple[str, str, str, str]] = set()
    counted_bible_exclusions: set[tuple[str, str, str, str]] = set()

    def process_candidate(brand_name: str, store: dict, matched_entries: list[dict]) -> None:
        nonlocal excluded_non_fss_fsf, excluded_in_bible, excluded_old_openings

        if not matched_entries:
            return

        candidate_key = _store_lookup_key(brand_name, store)
        if candidate_key in processed_candidates:
            return

        # A scraper's own door name and a research entry's door name can
        # both refer to the same physical BIBLE door under different text
        # (e.g. a locator's "Boutique Marais" vs the BIBLE's own "LE
        # MARAIS"). Check every name we have for this candidate - not just
        # whichever one built this particular store dict - so a door that's
        # already tracked under any of them is never re-reported as new.
        candidate_bible_keys = {candidate_key} | {
            _store_lookup_key(brand_name, _store_from_entry(entry)) for entry in matched_entries
        }
        matched_bible_key = next((key for key in candidate_bible_keys if key in bible_index["by_door_key"]), None)
        if matched_bible_key is not None:
            if matched_bible_key not in counted_bible_exclusions:
                excluded_in_bible += 1
                counted_bible_exclusions.add(matched_bible_key)
            processed_candidates.add(candidate_key)
            return

        sorted_entries = sorted(matched_entries, key=_recent_sort_key)
        recent_entry = next(
            (
                entry for entry in sorted_entries
                if entry.get("store_classification") in ALLOWED_DOOR_TYPES
                and _source_priority(entry) < 99
                and entry.get("status") in RECENT_OPENING_STATUSES
                and not _is_unattributed_linkedin_repost(entry)
                and _is_within_recent_days(_parse_source_date(entry.get("source_date")), checked_date, recent_days)
            ),
            None,
        )
        if recent_entry is not None:
            row = _recent_row(brand_name, store, recent_entry, matched_entries, checked_date)
            row_key = _row_key(row)
            if row_key not in recent_row_keys:
                recent_rows.append(row)
                recent_row_keys.add(row_key)
                recent_by_brand[brand_name] += 1
                if row["SOURCE URL"]:
                    source_links_used.add(row["SOURCE URL"])
            processed_candidates.add(candidate_key)
            return

        best_entry = _best_entry(matched_entries)
        if best_entry is None:
            return

        classification = best_entry.get("store_classification")

        if classification in ALLOWED_DOOR_TYPES:
            if _is_unattributed_linkedin_repost(best_entry):
                row = _verify_row(
                    brand_name, store, best_entry, matched_entries, checked_date,
                    "LinkedIn post without an identified original source", classification,
                )
                row_key = _row_key(row)
                if row_key not in verify_row_keys:
                    verify_rows.append(row)
                    verify_row_keys.add(row_key)
                    verify_by_brand[brand_name] += 1
                    if row["SOURCE URL"]:
                        source_links_used.add(row["SOURCE URL"])
                processed_candidates.add(candidate_key)
                return

            # TO VERIFY is only for candidates whose opening date is still
            # unknown. A known date that didn't qualify above (outside the
            # recent window, or reported by a source too weak to promote)
            # is a resolved fact, not an open question - it must never sit
            # in TO VERIFY. It simply isn't a recent opening to report.
            if any(_parse_source_date(e.get("source_date")) is not None for e in matched_entries):
                excluded_old_openings += 1
                processed_candidates.add(candidate_key)
                return

            reason = _candidate_reason(best_entry)
            if reason is None:
                return

            row = _verify_row(brand_name, store, best_entry, matched_entries, checked_date, reason, classification)
            row_key = _row_key(row)
            if row_key not in verify_row_keys:
                verify_rows.append(row)
                verify_row_keys.add(row_key)
                verify_by_brand[brand_name] += 1
                if row["SOURCE URL"]:
                    source_links_used.add(row["SOURCE URL"])
            processed_candidates.add(candidate_key)
            return

        if classification in OTHER_OPENING_CLASSIFICATIONS:
            row = _other_row(brand_name, store, best_entry, checked_date)
            row_key = _row_key(row)
            if row_key not in other_row_keys:
                other_rows.append(row)
                other_row_keys.add(row_key)
                other_by_brand[brand_name] += 1
                other_by_classification[classification] += 1
                if row["SOURCE URL"]:
                    source_links_used.add(row["SOURCE URL"])
            processed_candidates.add(candidate_key)
            return

        # PERFUMERY / MULTIBRAND_RETAILER / ONLINE / unrecognized - not the
        # brand's own doing at all (a reseller carrying the brand), pure
        # noise rather than any kind of opening to report.
        excluded_non_fss_fsf += 1
        processed_candidates.add(candidate_key)

    for brand_result in results:
        if brand_result.get("status") != "ok":
            continue
        brand_name = brand_result["brand"]

        for store in brand_result.get("stores", []):
            matched_entries = [
                entry for entry in entries
                if _entry_matches_store(entry, brand_name, store)
                and entry.get("status") in VERIFYABLE_STATUSES
            ]
            process_candidate(brand_name, store, matched_entries)

    for entry in entries:
        if entry.get("status") not in VERIFYABLE_STATUSES:
            continue
        brand_name = selected_brands[brand_match_key(entry.get("brand"))]
        process_candidate(brand_name, _store_from_entry(entry), [entry])

    summary = PipelineSummary(
        brands_checked=len(selected_brands),
        recent_openings_found=len(recent_rows),
        to_verify_count=len(verify_rows),
        excluded_non_fss_fsf=excluded_non_fss_fsf,
        excluded_in_bible=excluded_in_bible,
        excluded_old_openings=excluded_old_openings,
        other_openings_found=len(other_rows),
        excluded_travel_retail_gated=excluded_travel_retail_gated,
        excluded_industry_news_gated=excluded_industry_news_gated,
        recent_openings_by_brand=dict(sorted(recent_by_brand.items())),
        to_verify_by_brand=dict(sorted(verify_by_brand.items())),
        other_openings_by_brand=dict(sorted(other_by_brand.items())),
        other_openings_by_classification=dict(sorted(other_by_classification.items())),
        source_links_used=tuple(sorted(source_links_used)),
    )
    return recent_rows, verify_rows, other_rows, summary


def _write_sheet(workbook: openpyxl.Workbook, title: str, columns: list[str], rows: list[dict]) -> None:
    worksheet = workbook.create_sheet(title=title)
    worksheet.freeze_panes = "A2"
    worksheet.append(columns)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    for row in rows:
        worksheet.append([row.get(column) for column in columns])

    worksheet.auto_filter.ref = worksheet.dimensions

    if "SOURCE URL" in columns:
        hyperlink_col = columns.index("SOURCE URL") + 1
        for row_index in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_index, column=hyperlink_col)
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.style = "Hyperlink"

    if "SECOND SOURCE URL" in columns:
        hyperlink_col = columns.index("SECOND SOURCE URL") + 1
        for row_index in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_index, column=hyperlink_col)
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.style = "Hyperlink"

    for col_index, column in enumerate(columns, start=1):
        max_length = len(column)
        for row_index in range(2, worksheet.max_row + 1):
            value = worksheet.cell(row=row_index, column=col_index).value
            if value is None:
                continue
            max_length = max(max_length, len(str(value)))
        worksheet.column_dimensions[get_column_letter(col_index)].width = min(max_length + 2, 60)


def write_recent_openings_workbook(
    recent_rows: list[dict], verify_rows: list[dict], other_rows: list[dict], output_path: Path | str,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    _write_sheet(workbook, RECENT_OPENINGS_SHEET, RECENT_OPENINGS_COLUMNS, recent_rows)
    _write_sheet(workbook, TO_VERIFY_SHEET, TO_VERIFY_COLUMNS, verify_rows)
    _write_sheet(workbook, OTHER_OPENINGS_SHEET, OTHER_OPENINGS_COLUMNS, other_rows)
    workbook.save(output_path)
    workbook.close()
    return output_path
